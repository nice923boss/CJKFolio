"""docx 匯入（cjkfolio.docx_import）：OOXML 結構直接對映樣式（PLAN §4.4 修正 1）。

兩路 fixture：
- python-docx 產生器（Gate 4 獨立路徑：python-docx 寫、stdlib 讀）
- 手工 OOXML zip（邊界情境：val 關閉寫法、直掛 outlineLvl、語系名稱 fallback、壞檔）
"""
import zipfile
from pathlib import Path

import pytest

from cjkfolio.docx_import import import_docx, paragraphs_from_docx
from cjkfolio.styles import load_template

TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "neizai-zhengli.json"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


@pytest.fixture(scope="module")
def template():
    return load_template(TEMPLATE)


# ---------- 產生器路徑：python-docx 造檔，stdlib 讀 ----------

@pytest.fixture(scope="module")
def generated_docx(tmp_path_factory):
    import docx

    d = docx.Document()
    d.add_heading("第一章　整理的開始", level=1)
    d.add_heading("案例一　老公寓", level=2)
    d.add_heading("屋主的困境", level=3)
    p = d.add_paragraph("這間屋子住了三十年，")
    p.add_run("東西比人還多").bold = True
    p.add_run("，走道只剩一人寬。")
    p2 = d.add_paragraph("她說：")
    p2.add_run("捨不得丟").italic = True
    p2.add_run("，每一件都有故事。")
    d.add_paragraph("整理不是丟東西，是決定留下什麼。", style="Quote")
    d.add_paragraph("")  # 空段跳過
    both = d.add_paragraph("")
    r = both.add_run("粗斜並用")
    r.bold = True
    r.italic = True
    path = tmp_path_factory.mktemp("docx") / "chapter.docx"
    d.save(str(path))
    return path


def test_headings_detected_via_style_chain(generated_docx):
    kinds = [k for k, _ in paragraphs_from_docx(generated_docx)]
    assert kinds[:3] == ["h1", "h2", "h3"]


def test_quote_style_maps_blockquote(generated_docx):
    paras = paragraphs_from_docx(generated_docx)
    quote = [(k, rs) for k, rs in paras if k == "blockquote"]
    assert len(quote) == 1
    assert "".join(t for t, _, _ in quote[0][1]) == "整理不是丟東西，是決定留下什麼。"


def test_empty_paragraph_skipped(generated_docx):
    for _, runs in paragraphs_from_docx(generated_docx):
        assert "".join(t for t, _, _ in runs)  # 無空段


def test_body_runs_carry_bold_italic(generated_docx):
    paras = paragraphs_from_docx(generated_docx)
    body = [rs for k, rs in paras if k == "paragraph"]
    assert ("東西比人還多", True, False) in body[0]
    assert ("捨不得丟", False, True) in body[1]
    assert ("粗斜並用", True, True) in body[2]


def test_import_docx_maps_template_styles(generated_docx, template):
    doc = import_docx(generated_docx, mapping=template.md_mapping,
                      inline_mapping=template.md_inline_mapping, title="自擬稿")
    styles = [p.style for p in doc.paragraphs]
    assert styles[0] == "章名大標-明體W7"
    assert styles[1] == "個案大標-NotoSans"
    assert styles[2] == "章節小標-黑體"
    assert "引文-仿宋語錄" in styles
    assert doc.title == "自擬稿"
    assert doc.template == "neizai-zhengli"


def test_import_docx_inline_char_styles(generated_docx, template):
    doc = import_docx(generated_docx, mapping=template.md_mapping,
                      inline_mapping=template.md_inline_mapping)
    p = doc.paragraphs[3]
    assert p.text == "這間屋子住了三十年，東西比人還多，走道只剩一人寬。"
    assert [(r.text, r.char_style) for r in p.runs] == [
        ("這間屋子住了三十年，", None),
        ("東西比人還多", "內文強調-明體W9"),
        ("，走道只剩一人寬。", None),
    ]


def test_italic_maps_em_and_bold_italic_prefers_strong(generated_docx):
    doc = import_docx(generated_docx, mapping={"paragraph": "內文"},
                      inline_mapping={"strong": "粗樣式", "em": "斜樣式"})
    italic_para = doc.paragraphs[4]
    assert italic_para.runs[1].char_style == "斜樣式"
    both_para = doc.paragraphs[-1]
    assert both_para.runs[0].char_style == "粗樣式"  # 粗斜並用 → strong 優先


def test_title_defaults_to_file_stem(generated_docx):
    doc = import_docx(generated_docx, mapping={"paragraph": "內文"})
    assert doc.title == "chapter"


def test_mapping_requires_paragraph_key(generated_docx):
    with pytest.raises(ValueError, match="paragraph"):
        import_docx(generated_docx, mapping={})


# ---------- 手工 OOXML：邊界情境 ----------

DOC_TMPL = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W_NS}"><w:body>{{body}}</w:body></w:document>"""

STYLES_TMPL = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="{W_NS}">{{styles}}</w:styles>"""


def make_docx(tmp_path, body: str, styles: str = None, with_document=True) -> Path:
    path = tmp_path / "hand.docx"
    with zipfile.ZipFile(path, "w") as z:
        if with_document:
            z.writestr("word/document.xml", DOC_TMPL.format(body=body))
        if styles is not None:
            z.writestr("word/styles.xml", STYLES_TMPL.format(styles=styles))
    return path


def para(text: str, pstyle: str = None, outline: str = None, rpr: str = "") -> str:
    ppr = ""
    if pstyle or outline:
        ppr = "<w:pPr>"
        if pstyle:
            ppr += f'<w:pStyle w:val="{pstyle}"/>'
        if outline:
            ppr += f'<w:outlineLvl w:val="{outline}"/>'
        ppr += "</w:pPr>"
    rp = f"<w:rPr>{rpr}</w:rPr>" if rpr else ""
    return f"<w:p>{ppr}<w:r>{rp}<w:t>{text}</w:t></w:r></w:p>"


def style_def(sid: str, name: str, outline: str = None, based_on: str = None) -> str:
    parts = [f'<w:style w:type="paragraph" w:styleId="{sid}">', f'<w:name w:val="{name}"/>']
    if based_on:
        parts.append(f'<w:basedOn w:val="{based_on}"/>')
    if outline is not None:
        parts.append(f"<w:pPr><w:outlineLvl w:val=\"{outline}\"/></w:pPr>")
    parts.append("</w:style>")
    return "".join(parts)


def test_direct_outline_level_overrides_style(tmp_path):
    path = make_docx(
        tmp_path,
        para("直掛大綱", pstyle="Odd", outline="1"),
        styles=style_def("Odd", "Odd Style", outline="4"),
    )
    assert paragraphs_from_docx(path)[0][0] == "h2"  # 直掛 1 優先於樣式 4


def test_val_zero_and_false_disable_bold_italic(tmp_path):
    body = (
        para("關閉粗", rpr='<w:b w:val="0"/>')
        + para("關閉斜", rpr='<w:i w:val="false"/>')
        + para("開啟粗", rpr="<w:b/>")
    )
    paras = paragraphs_from_docx(make_docx(tmp_path, body))
    assert paras[0][1][0] == ("關閉粗", False, False)
    assert paras[1][1][0] == ("關閉斜", False, False)
    assert paras[2][1][0] == ("開啟粗", True, False)


def test_outline_level_nine_is_body(tmp_path):
    path = make_docx(
        tmp_path,
        para("目錄標題", pstyle="TOCHeading"),
        styles=style_def("TOCHeading", "TOC Heading", outline="9"),
    )
    assert paragraphs_from_docx(path)[0][0] == "paragraph"


def test_chinese_style_name_heading_fallback(tmp_path):
    path = make_docx(
        tmp_path,
        para("中文標題", pstyle="a2"),
        styles=style_def("a2", "標題 2"),  # 無 outlineLvl，靠名稱
    )
    assert paragraphs_from_docx(path)[0][0] == "h2"


def test_chinese_quote_name_maps_blockquote(tmp_path):
    path = make_docx(
        tmp_path,
        para("語錄", pstyle="yw"),
        styles=style_def("yw", "引文"),
    )
    assert paragraphs_from_docx(path)[0][0] == "blockquote"


def test_based_on_cycle_terminates(tmp_path):
    path = make_docx(
        tmp_path,
        para("循環", pstyle="A"),
        styles=style_def("A", "甲", based_on="B") + style_def("B", "乙", based_on="A"),
    )
    assert paragraphs_from_docx(path)[0][0] == "paragraph"


def test_missing_styles_xml_still_parses(tmp_path):
    path = make_docx(tmp_path, para("純內文", pstyle="Whatever"))
    assert paragraphs_from_docx(path) == [("paragraph", (("純內文", False, False),))]


def test_hyperlink_runs_included(tmp_path):
    body = (
        '<w:p><w:r><w:t>見</w:t></w:r>'
        '<w:hyperlink><w:r><w:t>連結文字</w:t></w:r></w:hyperlink>'
        "<w:r><w:t>。</w:t></w:r></w:p>"
    )
    paras = paragraphs_from_docx(make_docx(tmp_path, body))
    assert "".join(t for t, _, _ in paras[0][1]) == "見連結文字。"


def test_unmapped_h4_falls_back_to_paragraph_style(tmp_path):
    path = make_docx(tmp_path, para("深層標題", outline="3"))
    doc = import_docx(path, mapping={"paragraph": "內文-明體直排"})
    assert doc.paragraphs[0].style == "內文-明體直排"
    assert doc.paragraphs[0].text == "深層標題"


def test_adjacent_same_style_runs_merge(tmp_path):
    body = "<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>甲</w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>乙</w:t></w:r></w:p>"
    doc = import_docx(make_docx(tmp_path, body), mapping={"paragraph": "內文"},
                      inline_mapping={"strong": "強調"})
    assert [(r.text, r.char_style) for r in doc.paragraphs[0].runs] == [("甲乙", "強調")]


def test_cli_import_docx(generated_docx, tmp_path):
    from cjkfolio.cli import main
    from cjkfolio.document import load_document

    out = tmp_path / "chapter.doc.json"
    assert main(["import-docx", str(generated_docx), "-o", str(out), "--title", "自擬稿"]) == 0
    doc = load_document(out)
    assert doc.title == "自擬稿"
    assert doc.paragraphs[0].style == "章名大標-明體W7"
    assert any(r.char_style == "內文強調-明體W9" for p in doc.paragraphs for r in p.runs)


def test_not_a_zip_fast_fail(tmp_path):
    bogus = tmp_path / "fake.docx"
    bogus.write_text("這不是 zip", encoding="utf-8")
    with pytest.raises(ValueError, match="docx"):
        paragraphs_from_docx(bogus)


def test_zip_without_document_xml_fast_fail(tmp_path):
    path = make_docx(tmp_path, "", with_document=False, styles="")
    with pytest.raises(ValueError, match="document.xml"):
        paragraphs_from_docx(path)

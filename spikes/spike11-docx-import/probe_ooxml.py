"""Spike 11 探針：python-docx 產檔 → stdlib 解析 OOXML，驗證結構假設。

驗證項：
1. word/document.xml 段落序、pStyle、run 的 w:b／w:i（含 val="0" 關閉寫法）
2. word/styles.xml 的 styleId → name／outlineLvl／basedOn
3. 三層標題偵測鏈（段落直掛 outlineLvl → 樣式鏈 outlineLvl → 樣式名）
"""
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

OUT = Path(__file__).parent / "out"
OUT.mkdir(exist_ok=True)
DOCX = OUT / "probe.docx"

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def build_probe_docx() -> None:
    """獨立產生器：python-docx 造含標題／引文／粗斜體的 docx。"""
    import docx

    d = docx.Document()
    d.add_heading("第一章　整理的開始", level=1)
    d.add_heading("案例一　三十年的老公寓", level=2)
    d.add_heading("屋主的困境", level=3)
    p = d.add_paragraph("這間屋子住了三十年，")
    p.add_run("東西比人還多").bold = True
    p.add_run("，走道只剩一人寬。")
    p2 = d.add_paragraph("她說：")
    p2.add_run("捨不得丟").italic = True
    p2.add_run("，每一件都有故事。")
    d.add_paragraph("整理不是丟東西，是決定留下什麼。", style="Quote")
    d.add_paragraph("")  # 空段
    both = d.add_paragraph("")
    r = both.add_run("粗斜並用")
    r.bold = True
    r.italic = True
    d.save(str(DOCX))
    print(f"[產生器] python-docx {docx.__version__} → {DOCX.name}")


def read_xml(name: str) -> ET.Element:
    with zipfile.ZipFile(DOCX) as z:
        return ET.fromstring(z.read(name))


def flag_on(rpr, tag: str) -> bool:
    el = rpr.find(f"{W}{tag}") if rpr is not None else None
    if el is None:
        return False
    val = el.get(f"{W}val")
    return val not in ("0", "false", "none")


def main() -> None:
    build_probe_docx()

    with zipfile.ZipFile(DOCX) as z:
        names = z.namelist()
    print(f"[zip] {len(names)} 個成員；document.xml={'word/document.xml' in names} styles.xml={'word/styles.xml' in names}")

    # styles.xml：styleId → (name, outlineLvl, basedOn)
    styles = {}
    for st in read_xml("word/styles.xml").iter(f"{W}style"):
        sid = st.get(f"{W}styleId")
        name_el = st.find(f"{W}name")
        ppr = st.find(f"{W}pPr")
        lvl = None
        if ppr is not None:
            lvl_el = ppr.find(f"{W}outlineLvl")
            if lvl_el is not None:
                lvl = lvl_el.get(f"{W}val")
        based = st.find(f"{W}basedOn")
        styles[sid] = (
            name_el.get(f"{W}val") if name_el is not None else None,
            lvl,
            based.get(f"{W}val") if based is not None else None,
        )
    print("\n[styles.xml] styleId → (name, outlineLvl, basedOn)：")
    for sid, info in styles.items():
        if info[1] is not None or (info[0] or "").lower().startswith(("heading", "quote", "標題", "引文", "title", "normal")):
            print(f"  {sid!r} → {info}")

    # document.xml：逐段
    print("\n[document.xml] 逐段：")
    body = read_xml("word/document.xml").find(f"{W}body")
    for i, p in enumerate(body.findall(f"{W}p")):
        ppr = p.find(f"{W}pPr")
        pstyle = None
        direct_lvl = None
        if ppr is not None:
            ps = ppr.find(f"{W}pStyle")
            pstyle = ps.get(f"{W}val") if ps is not None else None
            ol = ppr.find(f"{W}outlineLvl")
            direct_lvl = ol.get(f"{W}val") if ol is not None else None
        runs = []
        for r in p.findall(f"{W}r"):
            rpr = r.find(f"{W}rPr")
            text = "".join(t.text or "" for t in r.findall(f"{W}t"))
            runs.append((text, flag_on(rpr, "b"), flag_on(rpr, "i")))
        # 三層偵測鏈
        lvl = direct_lvl
        seen = set()
        sid = pstyle
        while lvl is None and sid and sid in styles and sid not in seen:
            seen.add(sid)
            lvl = styles[sid][1]
            sid = styles[sid][2]
        sname = styles.get(pstyle, (None,))[0] if pstyle else None
        print(f"  段{i}: pStyle={pstyle!r} name={sname!r} 直掛lvl={direct_lvl!r} 鏈解lvl={lvl!r} runs={runs}")

    print("\n[結論] 三層鏈可行性見上；val 關閉寫法另以手工 XML 驗（測試層）。")


if __name__ == "__main__":
    main()
